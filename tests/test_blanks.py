"""Бумажные бланки: набор приборов, номер квартиры и разбивка по листам."""
import pytest
from openpyxl import load_workbook

from database import repository
from database.init_db import init_db
from database.models import apartment_meters
from excel.blanks import (DIGIT_CELLS, PAGE_BUDGET_PT, _blank_height,
                          generate_blanks)


@pytest.fixture()
def conn(tmp_path):
    db = tmp_path / "blanks.db"
    init_db(db, apartments_count=6, nonresidential_count=1)
    conn = repository.connect(db)
    # кв. 2 — трёхкомнатная: ХВС и ГВС раздельно по кухне и санузлу
    apt = repository.get_apartment_by_number(conn, "2")
    repository.set_meters(conn, apt["id"], apartment_meters(2, 2))
    conn.commit()
    yield conn
    conn.close()


def _labels(ws) -> list[str]:
    return [ws.cell(row=r, column=1).value
            for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=1).value]


def test_blank_per_flat_with_its_own_meters(conn, tmp_path):
    out = generate_blanks(conn, tmp_path / "blanki.xlsx", "август 2026")
    labels = _labels(load_workbook(out).active)

    # Номер квартиры напечатан заранее — перепутать его нельзя
    assert "Кв. 1" in labels and "Кв. 6" in labels
    # 1-2-комнатная: один ХВС и один ГВС
    assert "ХВС" in labels and "ГВС" in labels
    # 3-комнатная: раздельный учёт плюс сумма, которую ждёт ОЭК
    assert "ХВС кухня" in labels and "ГВС ванна" in labels
    assert any(str(v).startswith("Сумма ГВС") for v in labels)


def test_nonresidential_gets_no_blank(conn, tmp_path):
    """Нежилые и общедомовой прибор председатель передаёт сама."""
    out = generate_blanks(conn, tmp_path / "blanki.xlsx", "август 2026")
    labels = _labels(load_workbook(out).active)

    assert not any("Нежилое" in str(v) for v in labels)
    assert not any("Общедомовой" in str(v) for v in labels)


def test_only_requested_flats(conn, tmp_path):
    out = generate_blanks(conn, tmp_path / "blanki.xlsx", "август 2026",
                          numbers=["3", "5"])
    labels = _labels(load_workbook(out).active)

    assert [v for v in labels if str(v).startswith("Кв.")] == ["Кв. 3", "Кв. 5"]


def test_digit_cells_are_bordered(conn, tmp_path):
    """Клетка под цифру — та самая, ради которой всё затевалось."""
    out = generate_blanks(conn, tmp_path / "blanki.xlsx", "август 2026")
    ws = load_workbook(out).active

    row = next(r for r in range(1, ws.max_row + 1)
               if ws.cell(row=r, column=1).value == "Электроэнергия")
    for col in range(2, 2 + DIGIT_CELLS):
        assert ws.cell(row=row, column=col).border.left.style == "medium"


def test_blanks_are_not_split_between_pages(conn, tmp_path):
    """Разрыв листа только между бланками и не чаще, чем нужно."""
    out = generate_blanks(conn, tmp_path / "blanki.xlsx", "август 2026")
    ws = load_workbook(out).active

    heights = [_blank_height([m["kind"] for m
                              in repository.meters_for_apartment(conn, a["id"])])
               for a in repository.list_apartments(conn) if a["type"] == "residential"]

    pages, used = 1, 0.0
    for height in heights:
        if used and used + height > PAGE_BUDGET_PT:
            pages, used = pages + 1, 0.0
        used += height
    assert len(ws.row_breaks.brk) == pages - 1

    # Ни одна страница не выходит за высоту листа
    assert all(h <= PAGE_BUDGET_PT for h in heights)
