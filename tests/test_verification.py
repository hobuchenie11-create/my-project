"""Поверка общедомовых приборов: автоматический расчёт сроков."""
from datetime import date, timedelta

import pytest

from bot.services import verification_service as vs
from database import repository
from database.init_db import init_db
from database.models import VERIFICATION_LEAD_DAYS


@pytest.fixture()
def conn(tmp_path):
    db = tmp_path / "verify.db"
    init_db(db, apartments_count=1, nonresidential_count=0)
    conn = repository.connect(db)
    vs.ensure_house_meters(conn)
    yield conn
    conn.close()


def _meter(conn, part):
    return next(m for m in repository.house_meters(conn)
                if part.lower() in m["name"].lower())


def test_default_house_meters(conn):
    names = [m["name"] for m in repository.house_meters(conn)]
    assert any("Тепловая" in n for n in names)
    assert any("ГВС" in n for n in names)
    assert any("ХВС" in n for n in names)
    assert all(m["interval_years"] == 4 for m in repository.house_meters(conn))


def test_no_due_date_until_last_verification_entered(conn):
    meter = _meter(conn, "тепловая")
    v = vs.view(meter, date(2026, 9, 1))
    assert v.next_due is None
    assert v.mark == "⚪"
    assert "не внесена" in v.status_text


def test_next_due_is_computed_from_interval(conn):
    meter = _meter(conn, "тепловая")
    following = vs.register_verification(conn, meter["id"], date(2026, 3, 15))
    assert following == date(2030, 3, 15)        # 4 года

    meter = repository.get_house_meter(conn, meter["id"])
    assert meter["last_verified"] == "2026-03-15"
    assert vs.view(meter, date(2026, 9, 1)).next_due == date(2030, 3, 15)


def test_intervals_can_differ_per_meter(conn):
    """У каждого прибора свой межповерочный интервал."""
    heat = _meter(conn, "тепловая")
    cws = _meter(conn, "хвс")
    repository.update_house_meter(conn, heat["id"], interval_years=4)
    repository.update_house_meter(conn, cws["id"], interval_years=6)

    assert vs.register_verification(conn, heat["id"], date(2026, 5, 1)) == date(2030, 5, 1)
    assert vs.register_verification(conn, cws["id"], date(2026, 5, 1)) == date(2032, 5, 1)


def test_leap_day_does_not_break(conn):
    meter = _meter(conn, "гвс")
    # 29 февраля + 4 года: 2028 високосный, но проверим и невисокосный случай
    assert vs.add_years(date(2024, 2, 29), 1) == date(2025, 2, 28)
    assert vs.register_verification(conn, meter["id"], date(2028, 2, 29)) == date(2032, 2, 29)


def test_task_appears_only_within_lead_time(conn):
    meter = _meter(conn, "тепловая")
    vs.register_verification(conn, meter["id"], date(2026, 3, 15))  # срок 15.03.2030

    # За год до срока задачи ещё нет
    assert vs.sync_verification_tasks(conn, date(2029, 3, 15)) == 0
    # За полгода — появляется
    lead_day = date(2030, 3, 15) - timedelta(days=VERIFICATION_LEAD_DAYS)
    assert vs.sync_verification_tasks(conn, lead_day) == 1

    task = repository.verification_task(conn, meter["id"], "2030-03-15")
    assert task is not None
    assert task["category"] == "verification"
    assert task["priority"] == "high"
    assert task["start_date"] == lead_day.isoformat()


def test_task_is_not_duplicated(conn):
    meter = _meter(conn, "хвс")
    vs.register_verification(conn, meter["id"], date(2026, 1, 10))
    day = date(2030, 1, 10) - timedelta(days=10)
    assert vs.sync_verification_tasks(conn, day) == 1
    assert vs.sync_verification_tasks(conn, day) == 0


def test_new_verification_moves_the_next_date(conn):
    """После проведённой поверки срок пересчитывается на следующий цикл."""
    meter = _meter(conn, "гвс")
    vs.register_verification(conn, meter["id"], date(2026, 6, 1))
    assert vs.view(repository.get_house_meter(conn, meter["id"])).next_due \
        == date(2030, 6, 1)

    following = vs.register_verification(conn, meter["id"], date(2030, 5, 20))
    assert following == date(2034, 5, 20)

    history = repository.verification_history(conn, meter["id"])
    assert [h["verified_at"] for h in history] == ["2030-05-20", "2026-06-01"]


def test_overdue_verification(conn):
    meter = _meter(conn, "тепловая")
    vs.register_verification(conn, meter["id"], date(2020, 1, 1))  # срок 01.01.2024
    v = vs.view(repository.get_house_meter(conn, meter["id"]), date(2026, 9, 1))
    assert v.is_overdue
    assert v.mark == "🔴"
    assert "просрочена" in v.status_text


def test_verification_sheet_in_year_plan(conn, tmp_path):
    from excel.tasks_export import export_year_plan
    from openpyxl import load_workbook

    meter = _meter(conn, "тепловая")
    vs.register_verification(conn, meter["id"], date(2026, 3, 15))
    out = export_year_plan(conn, date.today().year, tmp_path / "plan.xlsx")

    wb = load_workbook(out)
    assert "Поверка приборов" in wb.sheetnames
    ws = wb["Поверка приборов"]
    assert ws.cell(2, 1).value == "Прибор учёта"
    names = [ws.cell(r, 1).value for r in range(3, 7)]
    assert any(n and "Тепловая" in n for n in names)
    row = next(r for r in range(3, 7) if "Тепловая" in (ws.cell(r, 1).value or ""))
    assert ws.cell(row, 3).value == "15.03.2026"     # последняя поверка
    assert ws.cell(row, 5).value == "15.03.2030"     # следующая — посчитана
