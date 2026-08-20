from openpyxl import load_workbook

from bot.services.reading_service import save_reading
from bot.services.report_service import STATEMENT_COLUMNS, build_statement
from database import repository
from database.init_db import init_db
from excel.export import PRINT_HEADERS, export_statement


def test_export_statement(tmp_path):
    db = tmp_path / "test.db"
    init_db(db, apartments_count=2, nonresidential_count=1)
    conn = repository.connect(db)
    try:
        apt = repository.get_apartment_by_number(conn, "1")
        save_reading(conn, apt["id"], "electricity", 1234, None, period="2026-07")
        statement = build_statement(conn, "2026-07")
    finally:
        conn.close()

    out = tmp_path / "vedomost.xlsx"
    export_statement(statement, "июль 2026", out)

    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Ведомость передачи показаний за июль 2026"
    assert [ws.cell(row=3, column=c).value
            for c in range(1, len(STATEMENT_COLUMNS) + 1)] == [
        PRINT_HEADERS.get(name, name) for name in STATEMENT_COLUMNS]
    # Нежилое помещение и общедомовой прибор идут первыми (первая страница),
    # затем квартиры по порядку. В печатной форме названия сокращены —
    # иначе узкая колонка «Кв.» рвёт их на три строки.
    assert ws.cell(row=4, column=1).value == "Нежилое №1"
    assert ws.cell(row=5, column=1).value == "Общедомовой ПУ"
    assert ws.cell(row=6, column=1).value == "1"
    assert ws.cell(row=6, column=2).value == "✔"
    assert ws.cell(row=6, column=3).value == 1234
    # Номер квартиры — по центру колонки: у края его неудобно читать в распечатке
    assert ws.cell(row=6, column=1).alignment.horizontal == "center"


def test_print_layout_splits_after_flat_40(tmp_path):
    """Печать: нежилые, ОДПУ и кв. 1–40 на первом листе, 41–80 на втором."""
    db = tmp_path / "print.db"
    init_db(db, apartments_count=80, nonresidential_count=2)
    conn = repository.connect(db)
    try:
        statement = build_statement(conn, "2026-07")
    finally:
        conn.close()

    out = tmp_path / "vedomost.xlsx"
    export_statement(statement, "июль 2026", out)

    ws = load_workbook(out).active
    # Разрыв листа стоит на строке с кв. 40
    breaks = [brk.id for brk in ws.row_breaks.brk]
    assert len(breaks) == 1
    assert ws.cell(row=breaks[0], column=1).value == "40"
    assert ws.cell(row=breaks[0] + 1, column=1).value == "41"
    # Шапка таблицы повторяется на втором листе — иначе непонятно, где какой
    # столбец, когда листы читают по отдельности
    assert ws.print_title_rows == "$3:$3"
