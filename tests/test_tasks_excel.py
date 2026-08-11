"""Круг «выгрузил → поправил в Excel → загрузил обратно».

Правки, сделанные в файле на компьютере, должны попадать в базу — иначе бот
о них не знает и продолжает напоминать про уже сделанное.
"""
from datetime import date

import pytest
from openpyxl import load_workbook

from bot.services import task_service, verification_service
from database import repository
from database.init_db import init_db
from excel.tasks_export import (COLUMNS, ONE_OFF_COLUMNS, SHEET_ONE_OFF,
                                SHEET_PLAN, SHEET_VERIFICATION,
                                VERIFICATION_COLUMNS, export_year_plan)
from excel.tasks_import import import_year_plan

YEAR = date.today().year


@pytest.fixture()
def conn(tmp_path):
    db = tmp_path / "tasks.db"
    init_db(db, apartments_count=2, nonresidential_count=1)
    conn = repository.connect(db)
    task_service.generate_year(conn, YEAR)
    verification_service.ensure_house_meters(conn)
    yield conn
    conn.close()


def _find_row(ws, id_column: int, value: int) -> int:
    for row in range(3, ws.max_row + 1):
        if ws.cell(row=row, column=id_column).value == value:
            return row
    raise AssertionError(f"строка с ID={value} не найдена")


def _by_title(conn, part: str):
    return next(r for r in repository.tasks_in_year(conn, YEAR)
                if part.lower() in r["title"].lower())


def test_plan_edits_return_to_the_database(conn, tmp_path):
    rent = _by_title(conn, "аренда за нежилое")
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")

    wb = load_workbook(path)
    ws = wb[SHEET_PLAN]
    row = _find_row(ws, len(COLUMNS), rent["id"])
    ws.cell(row=row, column=5, value="Выполнена")          # Статус
    ws.cell(row=row, column=6, value=35000)                # Аренда, ₽
    ws.cell(row=row, column=7, value="08.06.2026")         # Дата поступления
    ws.cell(row=row, column=10, value="Оплатил с опозданием")  # Комментарий
    wb.save(path)

    result = import_year_plan(conn, path, tg_id=1)
    assert result.updated == 1
    assert not result.problems

    updated = repository.get_task(conn, rent["id"])
    assert updated["status"] == "done"
    assert updated["amount"] == 35000
    assert updated["paid_at"] == "2026-06-08"
    assert updated["note"] == "Оплатил с опозданием"
    assert updated["done_at"]                       # дата выполнения проставлена

    history = repository.task_history(conn, rent["id"])
    assert any("Excel" in e["details"] for e in history)


def test_import_is_idempotent(conn, tmp_path):
    """Повторная загрузка того же файла ничего не меняет."""
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")
    first = import_year_plan(conn, path)
    second = import_year_plan(conn, path)
    assert (first.changed, second.changed) == (0, 0)


def test_empty_cell_does_not_erase_stored_value(conn, tmp_path):
    """Пустая ячейка — «не менять»: случайное стирание не удаляет данные."""
    rent = _by_title(conn, "аренда за нежилое")
    task_service.complete_task(conn, rent["id"], tg_id=1, amount=35000.0,
                               paid_at=f"{YEAR}-06-08", amount_field="amount")
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")

    wb = load_workbook(path)
    ws = wb[SHEET_PLAN]
    row = _find_row(ws, len(COLUMNS), rent["id"])
    ws.cell(row=row, column=6, value=None)
    ws.cell(row=row, column=7, value=None)
    wb.save(path)

    import_year_plan(conn, path)
    kept = repository.get_task(conn, rent["id"])
    assert kept["amount"] == 35000.0
    assert kept["paid_at"] == f"{YEAR}-06-08"


def test_new_row_on_my_tasks_becomes_a_task(conn, tmp_path):
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")

    wb = load_workbook(path)
    ws = wb[SHEET_ONE_OFF]
    row = 4                                   # ниже строки-подсказки
    ws.cell(row=row, column=1, value="Заказать смету на отмостку")
    ws.cell(row=row, column=2, value="Текущий ремонт")
    ws.cell(row=row, column=3, value=f"20.09.{YEAR}")
    wb.save(path)

    result = import_year_plan(conn, path, tg_id=1)
    assert result.created == 1

    task = repository.one_off_tasks(conn)[0]
    assert task["title"] == "Заказать смету на отмостку"
    assert task["category"] == "repair"
    assert task["due_date"] == f"{YEAR}-09-20"
    assert task["period"] == f"{YEAR}-09"

    # Второй раз строка уже с ID — дубля не будет
    again = import_year_plan(conn, export_year_plan(conn, YEAR,
                                                    tmp_path / "plan2.xlsx"))
    assert again.created == 0
    assert len(repository.one_off_tasks(conn)) == 1


def test_verification_dates_can_be_entered_in_excel(conn, tmp_path):
    """Даты поверки председатель вносит списком — сроки считаются сами."""
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")
    meter = repository.house_meters(conn)[0]

    wb = load_workbook(path)
    ws = wb[SHEET_VERIFICATION]
    row = _find_row(ws, len(VERIFICATION_COLUMNS), meter["id"])
    ws.cell(row=row, column=2, value="№ 12345")
    ws.cell(row=row, column=3, value="15.03.2024")
    ws.cell(row=row, column=4, value=6)
    wb.save(path)

    result = import_year_plan(conn, path)
    assert result.meters == 1

    updated = repository.get_house_meter(conn, meter["id"])
    assert updated["last_verified"] == "2024-03-15"
    assert updated["interval_years"] == 6
    assert updated["serial"] == "№ 12345"
    # Следующая поверка = 15.03.2024 + 6 лет
    assert verification_service.view(updated).next_due == date(2030, 3, 15)
    assert repository.verification_history(conn, meter["id"])


def test_unreadable_values_are_reported_not_swallowed(conn, tmp_path):
    rent = _by_title(conn, "аренда за нежилое")
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")

    wb = load_workbook(path)
    ws = wb[SHEET_PLAN]
    row = _find_row(ws, len(COLUMNS), rent["id"])
    ws.cell(row=row, column=5, value="сделано")            # нет такого статуса
    ws.cell(row=row, column=7, value="в понедельник")      # не дата
    wb.save(path)

    result = import_year_plan(conn, path)
    assert len(result.problems) == 2
    assert any("статус" in p for p in result.problems)
    assert any("дату" in p for p in result.problems)
    assert repository.get_task(conn, rent["id"])["status"] == "new"


def test_file_exported_before_the_id_column_still_loads(conn, tmp_path):
    """Старая выгрузка без «ID»: задача находится по месяцу и названию."""
    rent = _by_title(conn, "аренда за нежилое")
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")

    wb = load_workbook(path)
    ws = wb[SHEET_PLAN]
    row = _find_row(ws, len(COLUMNS), rent["id"])
    ws.cell(row=row, column=5, value="Выполнена")
    ws.cell(row=row, column=6, value=35000)
    ws.cell(row=row, column=7, value=f"08.06.{YEAR}")
    ws.delete_cols(len(COLUMNS))                  # файл вчерашней выгрузки
    ws.cell(row=2, column=10, value="Примечание")  # и старое имя столбца
    wb.save(path)

    result = import_year_plan(conn, path, tg_id=1)
    assert result.updated == 1
    assert not result.problems                    # итоги и заголовки — не ошибки

    updated = repository.get_task(conn, rent["id"])
    assert updated["status"] == "done"
    assert updated["amount"] == 35000
    assert updated["paid_at"] == f"{YEAR}-06-08"
    # Описание из регламента комментарием не считается
    assert updated["note"] == ""


def test_old_file_does_not_duplicate_my_tasks(conn, tmp_path):
    """Разовая задача из старого файла обновляется, а не задваивается."""
    repository.create_task(conn, "Заказать смету на отмостку", category="repair",
                           status="new", due_date=f"{YEAR}-09-20",
                           period=f"{YEAR}-09", source="chairman")
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")

    wb = load_workbook(path)
    ws = wb[SHEET_ONE_OFF]
    row = _find_row(ws, len(ONE_OFF_COLUMNS), repository.one_off_tasks(conn)[0]["id"])
    ws.cell(row=row, column=5, value="В работе")
    ws.delete_cols(len(ONE_OFF_COLUMNS))
    wb.save(path)

    result = import_year_plan(conn, path)
    assert (result.created, result.updated) == (0, 1)
    assert len(repository.one_off_tasks(conn)) == 1
    assert repository.one_off_tasks(conn)[0]["status"] == "in_progress"


def test_old_file_still_carries_verification_dates(conn, tmp_path):
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")
    meter = repository.house_meters(conn)[0]

    wb = load_workbook(path)
    ws = wb[SHEET_VERIFICATION]
    row = _find_row(ws, len(VERIFICATION_COLUMNS), meter["id"])
    ws.cell(row=row, column=3, value="15.03.2024")
    ws.delete_cols(len(VERIFICATION_COLUMNS))
    wb.save(path)

    assert import_year_plan(conn, path).meters == 1
    assert repository.get_house_meter(conn, meter["id"])["last_verified"] == "2024-03-15"


def test_foreign_workbook_is_rejected_with_a_hint(conn, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Лист1"
    other = tmp_path / "чужой.xlsx"
    wb.save(other)

    result = import_year_plan(conn, other)
    assert result.changed == 0
    assert "Годовой план" in result.problems[0]


def test_headers_survive_the_service_column(conn, tmp_path):
    """Столбец «ID» — служебный: он скрыт и стоит последним."""
    path = export_year_plan(conn, YEAR, tmp_path / "plan.xlsx")
    wb = load_workbook(path)

    for sheet, columns in ((SHEET_PLAN, COLUMNS),
                           (SHEET_ONE_OFF, ONE_OFF_COLUMNS),
                           (SHEET_VERIFICATION, VERIFICATION_COLUMNS)):
        ws = wb[sheet]
        assert ws.cell(row=2, column=len(columns)).value == "ID"
        letter = ws.cell(row=2, column=len(columns)).column_letter
        assert ws.column_dimensions[letter].hidden
