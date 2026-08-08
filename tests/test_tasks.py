"""Модуль «Задачи председателя»: годовой цикл, статусы, напоминания, сводки."""
from datetime import date

import pytest

from bot.services import task_service
from database import repository
from database.init_db import init_db


@pytest.fixture()
def conn(tmp_path):
    db = tmp_path / "tasks.db"
    init_db(db, apartments_count=2, nonresidential_count=1)
    conn = repository.connect(db)
    yield conn
    conn.close()


def _by_title(conn, period, part):
    return next(r for r in repository.tasks_for_period(conn, period)
                if part.lower() in r["title"].lower())


def test_templates_cover_the_chairman_cycle(conn):
    task_service.ensure_templates(conn)
    codes = {t["code"] for t in repository.active_task_templates(conn)}
    assert codes == {"bank_statement", "posting_invoices",
                     "nonresidential_payment", "nonresidential_utilities",
                     "submit_readings_rso"}


def test_generate_creates_tasks_with_windows(conn):
    created = task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    assert created == 5

    bank = _by_title(conn, "2026-09", "выписку из банка")
    assert bank["start_date"] == "2026-09-02"     # окно 2-5 числа
    assert bank["due_date"] == "2026-09-05"

    invoices = _by_title(conn, "2026-09", "квитанц")
    assert (invoices["start_date"], invoices["due_date"]) == ("2026-09-05", "2026-09-10")

    rent = _by_title(conn, "2026-09", "аренда за нежилое")
    assert rent["due_date"] == "2026-09-10"       # аренда — до 10 числа

    utilities = _by_title(conn, "2026-09", "коммунальных услуг")
    assert utilities["due_date"] == "2026-09-18"  # коммуналка — до 18 числа
    assert utilities["priority"] == "high"        # помечена как важная

    readings = _by_title(conn, "2026-09", "ресурсоснабжающим")
    assert (readings["start_date"], readings["due_date"]) == ("2026-09-20", "2026-09-25")


def test_generation_is_idempotent(conn):
    task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    again = task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    assert again == 0                            # повторный запуск не дублирует


def test_cycle_continues_into_next_year(conn):
    # Из декабря цикл разворачивается в январь-март следующего года
    task_service.generate_tasks(conn, date(2026, 12, 1), months_ahead=3)
    assert repository.tasks_for_period(conn, "2027-01")
    assert repository.tasks_for_period(conn, "2027-03")


def test_generate_year_covers_twelve_months(conn):
    created = task_service.generate_year(conn, 2027)
    assert created == 12 * 5
    assert len(repository.tasks_in_year(conn, 2027)) == 60


def test_short_month_does_not_break_dates(conn):
    # У шаблона день 25, февраль короче — дата остаётся корректной
    task_service.generate_tasks(conn, date(2027, 2, 1), months_ahead=0)
    readings = _by_title(conn, "2027-02", "ресурсоснабжающим")
    assert readings["due_date"] == "2027-02-25"


def test_overdue_and_active_states(conn):
    task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    bank = _by_title(conn, "2026-09", "выписку из банка")   # 2-5 сентября

    assert task_service.view(bank, date(2026, 9, 1)).is_active_now is False  # ещё рано
    assert task_service.view(bank, date(2026, 9, 3)).is_active_now is True
    assert task_service.view(bank, date(2026, 9, 10)).is_overdue is True


def test_rent_and_utilities_are_stored_separately(conn):
    """Аренда и коммуналка — разные задачи и разные поля сумм."""
    task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)

    rent = _by_title(conn, "2026-09", "аренда за нежилое")
    task_service.complete_task(conn, rent["id"], tg_id=1, amount=35000.0,
                               paid_at="2026-09-08", amount_field="amount")
    row = repository.get_task(conn, rent["id"])
    assert (row["amount"], row["paid_at"]) == (35000.0, "2026-09-08")
    assert row["utility_amount"] is None          # поле коммуналки не тронуто

    utilities = _by_title(conn, "2026-09", "коммунальных услуг")
    task_service.complete_task(conn, utilities["id"], tg_id=1, amount=4520.30,
                               paid_at="2026-09-17",
                               amount_field="utility_amount")
    row = repository.get_task(conn, utilities["id"])
    assert (row["utility_amount"], row["utility_paid_at"]) == (4520.30, "2026-09-17")
    assert row["amount"] is None                 # поле аренды не тронуто

    history = repository.task_history(conn, utilities["id"])
    assert any("коммуналка" in e["details"] for e in history)


def test_utilities_highlighted_three_days_before_due(conn):
    """Коммуналка до 18-го: с 15 числа задача считается горящей."""
    task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    utilities = _by_title(conn, "2026-09", "коммунальных услуг")

    assert task_service.view(utilities, date(2026, 9, 14)).is_soon is False
    assert task_service.view(utilities, date(2026, 9, 15)).is_soon is True
    assert task_service.view(utilities, date(2026, 9, 18)).is_soon is True
    assert task_service.view(utilities, date(2026, 9, 15)).mark == "🟠"

    messages = task_service.reminders_for_today(conn, date(2026, 9, 15))
    assert any("Скоро срок" in m and "коммунальных" in m for m in messages)
    assert any("❗" in m for m in messages)        # приоритет высокий


def test_one_off_task_counts_down_from_its_due_date(conn):
    """Разовая задача: после ввода срока бот сам считает, сколько осталось."""
    task_id = repository.create_task(
        conn, "Заказать смету на отмостку", category="repair",
        status="new", due_date="2026-09-20", period="2026-09", source="chairman")
    row = repository.get_task(conn, task_id)

    assert task_service.view(row, date(2026, 9, 18)).days_left == 2
    assert task_service.view(row, date(2026, 9, 18)).is_soon is True
    assert "осталось 2 дн." in task_service.task_line(row, date(2026, 9, 18))
    assert task_service.view(row, date(2026, 9, 25)).is_overdue is True
    assert "просрочено на 5 дн." in task_service.task_line(row, date(2026, 9, 25))

    # Разовые задачи видны отдельно от годового цикла
    one_off = repository.one_off_tasks(conn)
    assert [r["id"] for r in one_off] == [task_id]
    assert "Заказать смету" in task_service.one_off_text(conn, date(2026, 9, 18))


def test_done_task_is_not_reminded(conn):
    task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    bank = _by_title(conn, "2026-09", "выписку из банка")

    overdue_day = date(2026, 9, 12)
    assert any("выписку" in m for m in task_service.reminders_for_today(conn, overdue_day))

    task_service.set_status(conn, bank["id"], "done", tg_id=1)
    assert not any("выписку" in m
                   for m in task_service.reminders_for_today(conn, overdue_day))


def test_reminder_on_window_start(conn):
    task_service.generate_tasks(conn, date(2026, 9, 1), months_ahead=0)
    messages = task_service.reminders_for_today(conn, date(2026, 9, 2))
    assert any("Пора начинать" in m and "выписку" in m for m in messages)


def test_council_digest_hides_internal_details(conn):
    task_service.generate_tasks(conn, date.today(), months_ahead=0)
    period = task_service._month_period(date.today())
    payment = _by_title(conn, period, "коммунальных услуг")
    task_service.set_status(conn, payment["id"], "in_progress", tg_id=1)
    task_service.complete_task(conn, payment["id"], tg_id=1, amount=9999.0,
                               paid_at=date.today().isoformat(),
                               amount_field="utility_amount")

    digest = task_service.council_digest(conn)
    assert "Совет дома" in digest
    assert "9999" not in digest                  # суммы Совету не показываем
    assert "заполняется" not in digest.lower()


def test_year_plan_export(conn, tmp_path):
    from excel.tasks_export import export_year_plan
    from openpyxl import load_workbook

    task_service.generate_year(conn, 2027)
    out = export_year_plan(conn, 2027, tmp_path / "plan.xlsx")

    wb = load_workbook(out)
    assert wb.sheetnames == ["Годовой план", "Регламент"]
    ws = wb["Годовой план"]
    assert "2027" in ws.cell(1, 1).value
    # Четыре столбца по нежилому помещению: аренда и коммуналка
    header = [ws.cell(2, c).value for c in range(1, 11)]
    assert header[5:9] == ["Аренда, ₽", "Дата поступления",
                           "Оплата коммуналки, ₽", "Дата оплаты"]
    titles = [ws.cell(r, 2).value for r in range(3, 80)]
    assert any(t and "выписку из банка" in t.lower() for t in titles)
    assert any(t and "коммунальных услуг" in t.lower() for t in titles)
