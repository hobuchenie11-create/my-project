"""Реестр ОЭК: разбор шаблона ресурсника и заполнение колонки показаний."""
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
import xlrd
import xlwt
from openpyxl import Workbook, load_workbook

from excel.oek_registry import (DATE_FORMAT, OekFormatError, apartment_key,
                                data_rows, fill_registry, find_layout,
                                inspect_template, open_grids, sheet_role,
                                template_period)

# Шапка листа ОЭК: реквизиты договора, затем строка заголовков таблицы
HEADERS = ["Округ", "Квартира", "Наличие ПУ", "Номер ПУ",
           "Расход Э/Э за июль     2026, кВт*ч", "Метод расчета",
           "Дата показания", "Показания 07.26 - день", "Показания 07.26 - ночь",
           "Показания 07.26 - полупик", "Текущие показания- день",
           "Текущие показания- ночь", "Текущие показания- полупик",
           "Дата снятия показания"]
READING_COL = 10          # «Текущие показания- день» — колонка K
DATE_COL = 13             # «Дата снятия показания» — колонка N
FIRST_DATA_ROW = 8        # 0-based: строка 9 в Excel


def _preamble(service: str) -> list[list]:
    return [
        ["11.08.26 10:20"],
        ["Реестр индивидуальных приборов учета для внесения показаний  "],
        ["за август   2026"],
        ['УК: ООО "УК Евроцентр"'],
        ["№ договора: 55100001955436"],
        [f"Услуга: {service}"],
        ["Адрес электронной почты контактого лица:"],
    ]


def _rows(flats: list[str]) -> list[list]:
    return [["Советского АО физические лица ", flat, 1.0, f"71290270{flat:0>5}",
             100.0, "По показаниям", 46223.0, 1000.0, 0.0, 0.0, "", "", "", ""]
            for flat in flats]


def make_xls(path, flats=("1", "2", "3"), service="Электроэнергия",
             water_sheets=("Горячая вода", "Холодная вода")):
    """Книга в формате ОЭК: лист услуги плюс листы по воде."""
    book = xlwt.Workbook()
    date_style = xlwt.easyxf(num_format_str="m/d/yy")
    for name, service_name in [(service.upper(), service)] + [
            (w.upper(), w) for w in water_sheets]:
        sheet = book.add_sheet(name)
        grid = _preamble(service_name) + [HEADERS]
        grid += _rows(list(flats)) if service_name == service else []
        for r, row in enumerate(grid):
            for c, value in enumerate(row):
                if r >= FIRST_DATA_ROW and c == DATE_COL:
                    sheet.write(r, c, 46254.0, date_style)
                else:
                    sheet.write(r, c, value)
    book.save(str(path))
    return path


def make_xlsx(path, flats=("1", "2", "3")):
    book = Workbook()
    sheet = book.active
    sheet.title = "ЭЛЕКТРОЭНЕРГИЯ"
    for row in _preamble("Электроэнергия") + [HEADERS] + _rows(list(flats)):
        sheet.append(row)
    water = book.create_sheet("ГОРЯЧАЯ ВОДА")
    for row in _preamble("Горячая вода") + [HEADERS]:
        water.append(row)
    book.save(path)
    return path


# ---------------------------------------------------------------------------
# Разбор формы
# ---------------------------------------------------------------------------

def test_apartment_key_normalizes_registry_cells():
    assert apartment_key("7") == "7"
    assert apartment_key(7.0) == "7"
    assert apartment_key("007") == "7"
    assert apartment_key(" кв. 12 ") == "12"
    assert apartment_key("") == ""
    assert apartment_key("Итого") == ""


def test_sheet_role_reads_service_line(tmp_path):
    """«Адрес электронной почты» есть на каждом листе — по слову «электро»
    лист по воде когда-то считался электрическим."""
    grids = open_grids(make_xls(tmp_path / "oek.xls"))
    assert [sheet_role(g) for g in grids] == ["electricity", "water", "water"]


def test_find_layout_picks_day_column(tmp_path):
    """Из трёх зон («день», «ночь», «полупик») дом однотарифный — берём день."""
    grid = open_grids(make_xls(tmp_path / "oek.xls"))[0]
    layout = find_layout(grid)
    assert layout.header_row == FIRST_DATA_ROW - 1
    assert layout.apartment_col == 1
    assert layout.reading_col == READING_COL
    assert layout.date_col == DATE_COL
    assert [key for _, key in data_rows(grid, layout)] == ["1", "2", "3"]


def test_find_layout_survives_single_reading_column(tmp_path):
    """Ресурсник убрал зоны и оставил одну колонку — она и есть текущая."""
    headers = HEADERS[:READING_COL] + ["Текущие показания", "Дата снятия показания"]
    book = xlwt.Workbook()
    sheet = book.add_sheet("ЭЛЕКТРОЭНЕРГИЯ")
    for r, row in enumerate(_preamble("Электроэнергия") + [headers]):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    sheet.write(FIRST_DATA_ROW, 1, "5")
    path = tmp_path / "one.xls"
    book.save(str(path))

    layout = find_layout(open_grids(path)[0])
    assert layout.reading_col == READING_COL
    assert layout.date_col == READING_COL + 1


def test_unknown_form_raises(tmp_path):
    book = xlwt.Workbook()
    sheet = book.add_sheet("ЭЛЕКТРОЭНЕРГИЯ")
    sheet.write(0, 0, "Услуга: Электроэнергия")
    sheet.write(5, 0, "Что-то совсем другое")
    path = tmp_path / "broken.xls"
    book.save(str(path))

    try:
        find_layout(open_grids(path)[0])
    except OekFormatError as exc:
        assert "Квартира" in str(exc)
    else:
        raise AssertionError("ожидали OekFormatError")


def test_template_period_from_header(tmp_path):
    assert template_period(open_grids(make_xls(tmp_path / "oek.xls"))[0]) == "2026-08"


def test_inspect_template_reports_what_bot_found(tmp_path):
    info = inspect_template(make_xls(tmp_path / "oek.xls"))
    assert info.sheet == "ЭЛЕКТРОЭНЕРГИЯ"
    assert info.period == "2026-08"
    assert info.rows == 3
    assert info.reading_header == "Текущие показания- день"
    assert info.dropped_sheets == ["ГОРЯЧАЯ ВОДА", "ХОЛОДНАЯ ВОДА"]


# ---------------------------------------------------------------------------
# Заполнение
# ---------------------------------------------------------------------------

def test_fill_xls_writes_readings_and_drops_water(tmp_path):
    template = make_xls(tmp_path / "oek.xls", flats=("1", "2", "3"))
    out = tmp_path / "out.xls"
    result = fill_registry(template, {"1": 31668.0, "3": 12608.0}, out,
                           taken_on=date(2026, 8, 20),
                           known_apartments={"1", "2", "3"})

    assert result.filled == ["1", "3"]
    assert result.empty == ["2"]
    assert result.dropped_sheets == ["ГОРЯЧАЯ ВОДА", "ХОЛОДНАЯ ВОДА"]

    book = xlrd.open_workbook(out)
    assert book.sheet_names() == ["ЭЛЕКТРОЭНЕРГИЯ"]
    sheet = book.sheet_by_index(0)
    assert sheet.cell_value(FIRST_DATA_ROW, READING_COL) == 31668
    # Квартира без показаний остаётся пустой: ставить туда ноль нельзя —
    # ресурсник посчитает это показанием счётчика
    assert sheet.cell_value(FIRST_DATA_ROW + 1, READING_COL) == ""
    assert sheet.cell_value(FIRST_DATA_ROW + 2, READING_COL) == 12608
    # Дата снятия — настоящая дата, а не текст, и в привычном виде 20.08.2026:
    # у ОЭК в шаблоне стоит американское m/d/yy, читать его неудобно
    assert (xlrd.xldate.xldate_as_datetime(
        sheet.cell_value(FIRST_DATA_ROW, DATE_COL), book.datemode).date()
        == date(2026, 8, 20))
    formatted = xlrd.open_workbook(out, formatting_info=True)
    fmt = formatted.format_map[
        formatted.xf_list[formatted.sheet_by_index(0).cell_xf_index(
            FIRST_DATA_ROW, DATE_COL)].format_key].format_str
    assert fmt == DATE_FORMAT
    # Шаблон не тронут
    assert xlrd.open_workbook(template).sheet_by_index(0).cell_value(
        FIRST_DATA_ROW, READING_COL) == ""


def test_flat_without_a_reading_gets_no_date(tmp_path):
    """Нет показания — нет и даты снятия, даже если ресурсник её проставил.

    Иначе при загрузке в систему ОЭК строка выглядит снятой, хотя показания
    за квартиру никто не передавал.
    """
    template = make_xls(tmp_path / "oek.xls", flats=("1", "2"))
    # В шаблоне дата стоит у всех строк — её и нужно вычистить
    source = xlrd.open_workbook(template).sheet_by_index(0)
    assert source.cell_value(FIRST_DATA_ROW + 1, DATE_COL) != ""

    out = tmp_path / "out.xls"
    fill_registry(template, {"1": 100.0}, out, taken_on=date(2026, 8, 20),
                  known_apartments={"1", "2"})

    sheet = xlrd.open_workbook(out).sheet_by_index(0)
    assert sheet.cell_value(FIRST_DATA_ROW, DATE_COL) != ""       # кв. 1 сдала
    assert sheet.cell_value(FIRST_DATA_ROW + 1, DATE_COL) == ""   # кв. 2 — нет
    assert sheet.cell_value(FIRST_DATA_ROW + 1, READING_COL) == ""


def test_the_whole_date_column_uses_one_format(tmp_path):
    """Заполненные и пустые ячейки — в одном формате, колонка не пёстрая."""
    template = make_xls(tmp_path / "oek.xls", flats=("1", "2"))
    out = tmp_path / "out.xls"
    fill_registry(template, {"1": 100.0}, out, taken_on=date(2026, 8, 20))

    book = xlrd.open_workbook(out, formatting_info=True)
    sheet = book.sheet_by_index(0)
    formats = {book.format_map[
        book.xf_list[sheet.cell_xf_index(row, DATE_COL)].format_key].format_str
        for row in (FIRST_DATA_ROW, FIRST_DATA_ROW + 1)}
    assert formats == {DATE_FORMAT}


def test_fill_xls_keeps_template_formatting(tmp_path):
    """Оформление ресурсника трогать нельзя — реестр он читает глазами."""
    template = make_xls(tmp_path / "oek.xls")
    out = tmp_path / "out.xls"
    fill_registry(template, {"1": 100.0}, out)

    src = xlrd.open_workbook(template, formatting_info=True).sheet_by_index(0)
    dst = xlrd.open_workbook(out, formatting_info=True).sheet_by_index(0)
    # Всё, кроме заполняемых колонок, совпадает со шаблоном дословно
    for r in range(dst.nrows):
        for c in range(dst.ncols):
            if c in (READING_COL, DATE_COL):
                continue
            assert src.cell_value(r, c) == dst.cell_value(r, c), (r, c)
    assert sorted(src.merged_cells) == sorted(dst.merged_cells)
    assert (src.cell_xf_index(FIRST_DATA_ROW, READING_COL)
            == dst.cell_xf_index(FIRST_DATA_ROW, READING_COL))


def test_fill_reports_rows_outside_registry(tmp_path):
    """Нежилые и ОДПУ идут по другому договору — строк в реестре у них нет."""
    template = make_xls(tmp_path / "oek.xls", flats=("1", "2"))
    result = fill_registry(template,
                           {"1": 1.0, "Общедомовой прибор учета": 125087.0},
                           tmp_path / "out.xls",
                           known_apartments={"1", "2",
                                             "Общедомовой прибор учета"})
    assert result.not_in_registry == ["Общедомовой прибор учета"]
    assert result.unknown == []


def test_fill_separates_rows_missing_from_registry(tmp_path):
    """Квартира есть у ОЭК, но не в справочнике бота — это не «не передали»."""
    template = make_xls(tmp_path / "oek.xls", flats=("1", "99"))
    result = fill_registry(template, {"1": 1.0}, tmp_path / "out.xls",
                           known_apartments={"1"})
    assert result.empty == []
    assert result.unknown == ["99"]


def test_fill_xlsx_template(tmp_path):
    """Ресурсник может перейти на .xlsx — форму ищем так же, по заголовкам."""
    template = make_xlsx(tmp_path / "oek.xlsx", flats=("1", "2"))
    out = tmp_path / "out.xlsx"
    result = fill_registry(template, {"2": 555.0}, out,
                           taken_on=date(2026, 8, 20))

    assert result.filled == ["2"]
    assert result.dropped_sheets == ["ГОРЯЧАЯ ВОДА"]
    book = load_workbook(out)
    assert book.sheetnames == ["ЭЛЕКТРОЭНЕРГИЯ"]
    sheet = book["ЭЛЕКТРОЭНЕРГИЯ"]
    assert sheet.cell(row=FIRST_DATA_ROW + 2, column=READING_COL + 1).value == 555
    assert sheet.cell(row=FIRST_DATA_ROW + 1, column=READING_COL + 1).value in (None, "")
    taken_cell = sheet.cell(row=FIRST_DATA_ROW + 2, column=DATE_COL + 1)
    # openpyxl отдаёт дату как datetime — важно, что время нулевое
    assert taken_cell.value.date() == date(2026, 8, 20)
    assert (taken_cell.value.hour, taken_cell.value.minute) == (0, 0)
    assert taken_cell.number_format == DATE_FORMAT


# ---------------------------------------------------------------------------
# Сборка реестра целиком
# ---------------------------------------------------------------------------

@pytest.fixture()
def registry_env(tmp_path, monkeypatch):
    """Реестр в песочнице: своя база, своя папка шаблонов, свой выход."""
    from dataclasses import replace

    from bot.config import config
    from bot.services.reading_service import save_reading
    from database import repository
    from database.init_db import init_db
    from reports import oek_registry as reports_oek

    db = tmp_path / "oek.db"
    init_db(db, apartments_count=2, nonresidential_count=0)
    conn = repository.connect(db)
    try:
        apt = repository.get_apartment_by_number(conn, "1")
        save_reading(conn, apt["id"], "electricity", 31668, None,
                     period="2026-08")
    finally:
        conn.close()
    real_connect = repository.connect
    monkeypatch.setattr(repository, "connect",
                        lambda db_path=None: real_connect(db_path or db))

    templates = tmp_path / "tpl"
    templates.mkdir()
    make_xls(templates / "oek.xls", flats=("1", "2"))
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    monkeypatch.setattr(reports_oek, "config",
                        replace(config, reports_dir=out_dir, oek_dir=templates))
    return SimpleNamespace(module=reports_oek, out_dir=out_dir, db=db)


def test_taken_on_is_the_day_the_registry_was_built(registry_env):
    """«Дата снятия показания» — день обработки, в виде 20.08.2026."""
    result = registry_env.module.generate_oek_registry("2026-08")

    book = xlrd.open_workbook(result.path, formatting_info=True)
    sheet = book.sheet_by_index(0)
    written = xlrd.xldate.xldate_as_datetime(
        sheet.cell_value(FIRST_DATA_ROW, DATE_COL), book.datemode)
    assert written.date() == date.today()
    assert (written.hour, written.minute) == (0, 0)      # время не пишем
    fmt = book.format_map[
        book.xf_list[sheet.cell_xf_index(FIRST_DATA_ROW, DATE_COL)].format_key]
    assert fmt.format_str == DATE_FORMAT


def test_locked_output_file_does_not_stop_the_registry(registry_env, monkeypatch):
    """Прошлый реестр открыт в Excel — новый должен лечь рядом, а не пропасть.

    Windows не даёт перезаписать открытый файл, а держать реестр открытым —
    обычное дело: 20 числа в 14:30 выгрузка из-за этого срывалась целиком.
    """
    reports_oek = registry_env.module
    out_dir = registry_env.out_dir

    # Занятое имя: обращение к нему падает так же, как у Windows
    taken = out_dir / "reestr_oek_2026-08.xls"
    taken.write_bytes(b"open in Excel")
    real_fill = reports_oek.fill_registry

    def fill(template, readings, path, **kwargs):
        if Path(path) == taken:
            raise PermissionError(13, "Permission denied")
        return real_fill(template, readings, path, **kwargs)

    monkeypatch.setattr(reports_oek, "fill_registry", fill)

    result = reports_oek.generate_oek_registry("2026-08")

    assert result.path != taken
    assert result.path.exists()
    assert result.path.name.startswith("reestr_oek_2026-08_")
    assert taken.read_bytes() == b"open in Excel"   # открытый файл не тронут


# Когда реестр уходит — в tests/test_scheduler.py, рядом с остальным календарём


def test_summary_lists_who_did_not_report(tmp_path):
    template = make_xls(tmp_path / "oek.xls", flats=("1", "2", "3"))
    result = fill_registry(template, {"1": 1.0}, tmp_path / "out.xls",
                           known_apartments={"1", "2", "3"})
    text = result.summary("август 2026")
    assert "август 2026" in text
    assert "Заполнено показаний: 1 из 3" in text
    assert "Не передали: 2, 3" in text
