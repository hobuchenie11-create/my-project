"""Книга Excel «Сбор показаний»: состав листов и наполнение."""
import pytest
from openpyxl import load_workbook

from bot.services.parser import parse_message
from bot.services.reading_service import save_parsed_readings
from database import repository
from database.init_db import init_db
from database.models import apartment_meters
from excel import style
from excel.workbook import (REGISTRY_COLUMNS, SHEET_CONTROL, SHEET_CURRENT,
                            SHEET_HISTORY, SHEET_REGISTRY, SHEET_SETTINGS,
                            build_workbook)


@pytest.fixture()
def conn(tmp_path):
    db = tmp_path / "test.db"
    init_db(db, apartments_count=3, nonresidential_count=0)
    conn = repository.connect(db)
    # кв.1 — 3-комнатная с раздельным учетом, кв.2 — 1-комнатная
    apt1 = repository.get_apartment_by_number(conn, "1")
    repository.upsert_apartment(conn, "1", "residential", 1, rooms=3)
    repository.set_meters(conn, apt1["id"], apartment_meters(2, 2))
    repository.upsert_apartment(conn, "2", "residential", 2, rooms=1)
    conn.commit()
    yield conn
    conn.close()


def _build(conn, tmp_path):
    out = tmp_path / "book.xlsx"
    build_workbook(conn, out, "2026-07")
    return load_workbook(out)


def test_sheets_present(conn, tmp_path):
    wb = _build(conn, tmp_path)
    assert wb.sheetnames == [SHEET_REGISTRY, SHEET_HISTORY, SHEET_CURRENT,
                             SHEET_CONTROL, SHEET_SETTINGS]


def test_registry_columns_and_rows(conn, tmp_path):
    apt = repository.get_apartment_by_number(conn, "1")
    repository.create_user(conn, 555, "Иванова М.П.", apt["id"], username="maria")
    parsed = parse_message("Кв 1\nЭл 15230\nХвс кухня 120\nХвс санузел 45\n"
                           "Гвс кухня 60\nГвс ванна 30")
    save_parsed_readings(conn, apt, parsed, None, period="2026-07")

    ws = _build(conn, tmp_path)[SHEET_REGISTRY]
    assert [ws.cell(2, c).value for c in range(1, len(REGISTRY_COLUMNS) + 1)] \
        == REGISTRY_COLUMNS

    row = [ws.cell(3, c).value for c in range(1, 18)]
    assert row[0] == "1"
    assert row[1] == "3-комнатная"
    assert row[2] == 555                       # Telegram ID
    assert row[4:9] == [120, 45, 60, 30, 15230]
    assert row[10] == style.STATUS_SUBMITTED
    assert row[12] == "@maria"
    assert ws.freeze_panes == "C3"
    assert ws.auto_filter.ref is not None


def test_status_no_telegram_and_not_submitted(conn, tmp_path):
    apt = repository.get_apartment_by_number(conn, "2")
    repository.create_user(conn, 777, "Петров И.И.", apt["id"])
    ws = _build(conn, tmp_path)[SHEET_REGISTRY]
    statuses = {ws.cell(r, 1).value: ws.cell(r, 11).value for r in range(3, 6)}
    assert statuses["2"] == style.STATUS_NOT_SUBMITTED  # есть Telegram, не сдал
    assert statuses["3"] == style.STATUS_NO_TELEGRAM    # не зарегистрирован


def test_submitted_status_without_registration(conn, tmp_path):
    # Факт передачи важнее регистрации
    apt = repository.get_apartment_by_number(conn, "2")
    parsed = parse_message("Кв 2\nЭл 100\nХвс 5\nГвс 7")
    save_parsed_readings(conn, apt, parsed, None, period="2026-07")
    ws = _build(conn, tmp_path)[SHEET_REGISTRY]
    row = next(r for r in range(3, 6) if ws.cell(r, 1).value == "2")
    assert ws.cell(row, 11).value == style.STATUS_SUBMITTED


def test_history_and_current_sheets(conn, tmp_path):
    apt = repository.get_apartment_by_number(conn, "2")
    parsed = parse_message("Кв 2\nЭл.эн 8204\nХвс 15\nГвс 22")
    save_parsed_readings(conn, apt, parsed, None, source="bot", period="2026-07")

    wb = _build(conn, tmp_path)
    hist = wb[SHEET_HISTORY]
    assert hist.cell(1, 1).value == "Дата"
    assert hist.cell(2, 3).value == "2"
    assert hist.cell(2, 9).value == "Telegram (бот)"

    cur = wb[SHEET_CURRENT]
    row = next(r for r in range(2, 6) if cur.cell(r, 1).value == "2")
    assert cur.cell(row, 3).value == 15       # ХВС кухня (единственный ХВС)
    assert cur.cell(row, 7).value == 8204     # Электроэнергия


def test_control_and_settings(conn, tmp_path):
    wb = _build(conn, tmp_path)
    control = wb[SHEET_CONTROL]
    labels = [control.cell(r, 1).value for r in range(3, 10)]
    assert "Всего квартир" in labels
    assert "Процент передачи" in labels
    assert any(str(control.cell(r, 2).value or "").startswith("=")
               for r in range(3, 10))         # формулы, а не статичные числа

    settings = wb[SHEET_SETTINGS]
    assert settings.cell(2, 1).value == "Статусы"
    assert settings.cell(3, 1).value == style.STATUS_SUBMITTED
