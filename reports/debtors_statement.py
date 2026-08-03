"""Ведомость непередавших показания (печатная форма).

Формируется по кнопке председателя и автоматически — в назначенный день
и час (по умолчанию 20 число, 09:00; см. DEBTORS_DAY/DEBTORS_HOUR в .env).

Вручную:  python -m reports.debtors_statement [ГГГГ-ММ]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from bot.config import config
from bot.services.reading_service import current_period, period_title
from database import repository
from excel import style

COLUMNS = ["№ квартиры", "Тип", "Житель", "Telegram", "Последняя передача", "Отметка"]
WIDTHS = [14, 18, 30, 18, 20, 14]


def build_debtors_statement(period: str, out_path: Path,
                            db_path: Path | str | None = None) -> tuple[Path, int]:
    """Собирает ведомость непередавших. Возвращает путь и число должников."""
    conn = repository.connect(db_path)
    try:
        debtors = repository.debtors(conn, period)
        registry = {r["number"]: r for r in repository.registry_rows(conn)}
        total = len(repository.list_apartments(conn))
    finally:
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Не передали"

    ncols = len(COLUMNS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(row=1, column=1,
                    value=f"Ведомость непередавших показания — {period_title(period)}")
    title.font = style.FONT_TITLE
    title.fill = style.FILL_TITLE
    title.alignment = style.CENTER
    ws.row_dimensions[1].height = 24

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    subtitle = ws.cell(row=2, column=1,
                       value=f"Не передали: {len(debtors)} из {total} помещений")
    subtitle.alignment = style.CENTER

    for col, name in enumerate(COLUMNS, start=1):
        ws.cell(row=3, column=col, value=name)
    style.style_header(ws, 3, ncols, WIDTHS)

    row = 4
    for item in debtors:
        info = registry.get(item["number"])
        rooms = info["rooms"] if info else 0
        type_ = info["type"] if info else "residential"
        username = info["username"] if info else ""
        cells = [
            item["number"],
            style.room_label(rooms, type_),
            item["full_name"] or "не зарегистрирован",
            f"@{username}" if username else ("есть" if item["tg_id"] else "нет"),
            _short(info["last_submission"] if info else None),
            "",  # отметка председателя от руки
        ]
        for col, value in enumerate(cells, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = style.BORDER
            cell.alignment = style.LEFT if col == 3 else style.CENTER
            if col == 1:
                cell.fill = (style.FILL_NO_TELEGRAM if not item["tg_id"]
                             else style.FILL_NOT_SUBMITTED)
        row += 1

    ws.cell(row=row + 1, column=1,
            value="🩶 серым — квартиры без регистрации в Telegram")
    ws.cell(row=row + 3, column=1, value="Председатель: ______________________")

    _setup_print(ws, last_row=row + 3, ncols=ncols)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, len(debtors)


def _setup_print(ws, last_row: int, ncols: int) -> None:
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    ws.print_area = f"A1:{get_column_letter(ncols)}{last_row}"
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.freeze_panes = "A4"
    ws.oddFooter.right.text = "Стр. &P из &N"


def _short(value: str | None) -> str:
    if not value:
        return "—"
    try:
        day, time = value.split(" ")
        y, m, d = day.split("-")
        return f"{d}.{m}.{y} {time[:5]}"
    except (ValueError, AttributeError):
        return value


def generate_debtors_statement(period: str | None = None) -> tuple[Path, int]:
    period = period or current_period()
    out = config.reports_dir / f"ne_peredali_{period}.xlsx"
    return build_debtors_statement(period, out)


if __name__ == "__main__":
    path, count = generate_debtors_statement(sys.argv[1] if len(sys.argv) > 1 else None)
    print(f"Ведомость непередавших сохранена: {path} (должников: {count})")
