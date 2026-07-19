# -*- coding: utf-8 -*-
"""
Демо-версия пакета документов для показа потенциальному клиенту.

Берёт базовую книгу (собирает её скриптом build_lift_docs.py), добавляет
титульный лист «Демо» со сценарием демонстрации и блоком контактов,
проставляет «ДЕМО-ВЕРСИЯ» в печатные колонтитулы всех листов.
"""
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = "Пакет_документов_монтаж_лифтов.xlsx"
OUT = "Пакет_документов_монтаж_лифтов_ДЕМО.xlsx"

ROOT = Path(__file__).resolve().parent.parent
subprocess.run([sys.executable, str(ROOT / "scripts" / "build_lift_docs.py")], check=True, cwd=ROOT)

wb = load_workbook(ROOT / BASE)

F = "Arial"
FONT = Font(name=F, size=10)
FONT_B = Font(name=F, size=10, bold=True)
FONT_TITLE = Font(name=F, size=15, bold=True, color="1F4E79")
FONT_H = Font(name=F, size=11, bold=True, color="1F4E79")
FONT_STEP = Font(name=F, size=10, bold=True)
FONT_INPUT = Font(name=F, size=10, color="0000FF")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_BAND = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="9DC3E6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AL = Alignment(horizontal="left", vertical="top", wrap_text=True)

ws = wb.create_sheet("Демо", 0)
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 105
ws.column_dimensions["C"].width = 3
ws.sheet_view.showGridLines = False


def row(r, text, font=FONT, height=None, fill=None, border=None):
    c = ws.cell(row=r, column=2, value=text)
    c.font = font
    c.alignment = AL
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if height:
        ws.row_dimensions[r].height = height
    return c


r = 2
row(r, "СИСТЕМА ПОДГОТОВКИ ПАКЕТА ДОКУМЕНТОВ ПО МОНТАЖУ ЛИФТОВ", FONT_TITLE, height=22)
r += 1
row(r, "ДЕМО-ВЕРСИЯ. Все организации, реквизиты, объёмы и цены в файле вымышленные.",
    Font(name=F, size=10, italic=True, color="C00000"), height=16)
r += 2

row(r, "ЧТО ЭТО", FONT_H)
r += 1
row(r, "Один файл — девять взаимосвязанных документов: Параметры → Смета → Спецификация → "
       "КП → КС-2 → КС-3 → Счёт → Сопроводительное письмо → Договор. Все реквизиты сторон, объект, "
       "номера и даты вводятся один раз на листе «Параметры», позиции работ — один раз на листе «Смета». "
       "Остальные документы формируются автоматически: суммы, НДС, номера, даты и даже текст договора "
       "пересчитываются синхронно. Ничего не «расходится» между сметой, актами и счётом.", height=68)
r += 2

row(r, "ПОПРОБУЙТЕ САМИ — 3 ШАГА (правьте только жёлтые ячейки)", FONT_H)
r += 1
steps = [
    ("Шаг 1. Смените номер договора.",
     "Лист «Параметры», ячейка B5 — впишите любой номер. Откройте листы «Договор», «КС-2», «Счёт», "
     "«Сопроводительное письмо»: номер обновился везде, включая заголовок договора и тему письма."),
    ("Шаг 2. Измените цену любой позиции сметы.",
     "Лист «Смета», жёлтая колонка «Цена за ед.» — поменяйте цифру. Итоги сметы, КП, акт КС-2, справка "
     "КС-3, счёт и пункт 2.1 договора пересчитаются мгновенно."),
    ("Шаг 3. Поменяйте ставку НДС.",
     "Лист «Параметры», ячейка B12 (например, 0% или 5%). Строки НДС и суммы «с НДС» обновятся во всех "
     "документах, включая текст договора."),
]
for title, body in steps:
    row(r, title, FONT_STEP, fill=FILL_BAND, border=BORDER)
    r += 1
    row(r, body, FONT, height=40, border=BORDER)
    r += 1
r += 1

row(r, "ЧТО НАСТРАИВАЕТСЯ ПОД ВАШУ ОРГАНИЗАЦИЮ", FONT_H)
r += 1
row(r, "• Ваши реквизиты, подписанты и банковские данные — из карточки предприятия.\n"
       "• Ваши привычные формулировки договора, порядок оплаты, гарантийные условия.\n"
       "• Печатные формы КС-2, КС-3 и счёта — под требования вашей бухгалтерии и заказчиков.\n"
       "• Импорт смет из вашей сметной программы (Гранд-Смета, Smeta.ru, 1С и др.): новая смета "
       "заводится одной вставкой, без ручного набора.", height=68)
r += 2

row(r, "КОНТАКТЫ", FONT_H)
r += 1
contacts = [("Исполнитель:", "впишите ваше имя"),
            ("Телефон:", "впишите телефон"),
            ("E-mail:", "h.obuchenie11@gmail.com")]
for label, value in contacts:
    ws.cell(row=r, column=2, value=label).font = FONT_B
    c = ws.cell(row=r, column=3, value=value)
    c.font = FONT_INPUT
    c.fill = FILL_INPUT
    c.alignment = AL
    ws.column_dimensions["C"].width = 40
    r += 1

ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1

for name in wb.sheetnames:
    sh = wb[name]
    sh.oddHeader.center.text = "ДЕМО-ВЕРСИЯ — данные вымышленные"
    sh.oddHeader.center.size = 8
    sh.oddHeader.center.font = "Arial,Italic"

wb.save(ROOT / OUT)
print("saved", OUT)
