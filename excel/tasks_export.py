"""Годовой план задач председателя в Excel.

Лист «Годовой план» — все задачи года по месяцам: срок, статус, сумма и дата
оплаты. Лист «Регламент» — расшифровка регулярных задач и их окон, чтобы
план был понятен и через год, и следующему председателю.

Файл рассчитан на обратную загрузку: у каждой строки есть скрытый столбец
«ID» — по нему `excel.tasks_import` находит задачу в базе и переносит в неё
правки, сделанные в Excel. Столбец не трогаем и не удаляем.
"""
import sqlite3
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from bot.config import config
from bot.services import task_service
from database import repository
from database.models import DEFAULT_TASK_TEMPLATES, TASK_CATEGORIES, TASK_STATUSES
from excel import style

SHEET_PLAN = "Годовой план"
SHEET_ONE_OFF = "Мои задачи"
SHEET_VERIFICATION = "Поверка приборов"
SHEET_RULES = "Регламент"

# Служебный столбец: связывает строку таблицы с записью в базе
COL_ID_TITLE = "ID"

# Пометка строк-подсказок: при обратной загрузке они пропускаются
HINT_MARK = "ℹ️"

# «Вид» — только для чтения: во вкладке рядом с приборами учёта живёт
# оборудование с гарантией (лифт), и без этого столбца его строка читалась
# бы как поверка. Названия остальных столбцов менять нельзя — по ним
# обратная загрузка находит данные в уже выгруженных файлах.
VERIFICATION_COLUMNS = ["Прибор учёта", "Вид", "Заводской №",
                        "Последняя поверка", "Интервал, лет",
                        "Следующая поверка", "Осталось",
                        "Примечание", COL_ID_TITLE]
VERIFICATION_WIDTHS = [38, 14, 18, 18, 14, 20, 22, 30, 6]

# Лист «Мои задачи» — разовые дела председателя, вне регулярного цикла
ONE_OFF_COLUMNS = ["Задача", "Категория", "Срок", "Осталось", "Статус",
                   "Создана", "Выполнена", "Примечание", COL_ID_TITLE]
ONE_OFF_WIDTHS = [44, 24, 13, 16, 14, 13, 13, 34, 6]

COLUMNS = ["Месяц", "Задача", "Категория", "Срок", "Статус",
           "Аренда, ₽", "Дата поступления",
           "Оплата коммуналки, ₽", "Дата оплаты", "Комментарий", COL_ID_TITLE]
WIDTHS = [14, 42, 20, 12, 14, 13, 17, 20, 14, 30, 6]

# Столбцы с суммами (для формата и итогов)
COL_RENT, COL_RENT_DATE, COL_UTIL, COL_UTIL_DATE = 6, 7, 8, 9

def _list_validation(values) -> DataValidation:
    """Выпадающий список: правки в Excel возвращаются понятными значениями."""
    return DataValidation(type="list", allow_blank=True,
                          formula1='"' + ",".join(values) + '"')


def _hide_service_column(ws, column: int) -> None:
    """Прячет служебный столбец «ID» — он нужен загрузке, а не глазам."""
    ws.column_dimensions[get_column_letter(column)].hidden = True

FILL_DONE = PatternFill("solid", fgColor="D9EAD3")      # выполнено
FILL_OVERDUE = PatternFill("solid", fgColor="F4CCCC")   # просрочено
FILL_ACTIVE = PatternFill("solid", fgColor="FFF2CC")    # в работе сейчас
FILL_SOON = PatternFill("solid", fgColor="FCE5CD")      # срок через 3 дня и меньше
FILL_MONTH = PatternFill("solid", fgColor="EFEFEF")


def export_year_plan(conn: sqlite3.Connection, year: int,
                     out_path: Path | None = None) -> Path:
    out_path = out_path or config.reports_dir / f"plan_zadach_{year}.xlsx"
    wb = Workbook()
    _sheet_plan(wb.active, conn, year)
    _sheet_one_off(wb.create_sheet(SHEET_ONE_OFF), conn)
    _sheet_verification(wb.create_sheet(SHEET_VERIFICATION), conn)
    _sheet_rules(wb.create_sheet(SHEET_RULES))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _sheet_plan(ws, conn: sqlite3.Connection, year: int) -> None:
    ws.title = SHEET_PLAN
    ncols = len(COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(row=1, column=1, value=f"План задач председателя на {year} год")
    title.font = style.FONT_TITLE
    title.fill = style.FILL_TITLE
    title.alignment = style.CENTER
    ws.row_dimensions[1].height = 24

    for col, name in enumerate(COLUMNS, start=1):
        ws.cell(row=2, column=col, value=name)
    style.style_header(ws, 2, ncols, WIDTHS)

    today = date.today()
    rows = repository.tasks_in_year(conn, year)
    by_period: dict[str, list] = {}
    for row in rows:
        key = row["period"] or (row["due_date"][:7] if row["due_date"] else "—")
        by_period.setdefault(key, []).append(row)

    r = 3
    for period in sorted(by_period):
        label = (task_service.period_title(period) if "-" in period
                 else "Без периода")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        month_cell = ws.cell(row=r, column=1, value=label.capitalize())
        month_cell.font = style.FONT_BOLD
        month_cell.fill = FILL_MONTH
        month_cell.alignment = style.LEFT
        r += 1

        for row in by_period[period]:
            v = task_service.view(row, today)
            cells = [
                "",
                row["title"],
                task_service.category_label(row["category"]),
                _fmt(row["due_date"]),
                task_service.status_label(row["status"]),
                row["amount"] if row["amount"] is not None else "",
                _fmt(row["paid_at"]),
                row["utility_amount"] if row["utility_amount"] is not None else "",
                _fmt(row["utility_paid_at"]),
                row["note"],          # свободный комментарий, правится в Excel
                row["id"],
            ]
            fill = (FILL_DONE if row["status"] == "done"
                    else FILL_OVERDUE if v.is_overdue
                    else FILL_SOON if v.is_soon
                    else FILL_ACTIVE if v.is_active_now else None)
            for col, value in enumerate(cells, start=1):
                cell = ws.cell(row=r, column=col, value=value)
                cell.border = style.BORDER
                cell.alignment = (style.LEFT if col in (2, 3, 10) else style.CENTER)
                # Подсветкой отмечаем задачу целиком: срок, статус и суммы
                if fill and col in (4, 5, COL_RENT, COL_UTIL):
                    cell.fill = fill
                if col in (COL_RENT, COL_UTIL):
                    cell.number_format = "# ##0.00"
                if row["priority"] == "high" and col == 2:
                    cell.font = style.FONT_BOLD
            r += 1

    last_data_row = r - 1

    # Итоги по суммам оплат за год
    r += 1
    ws.cell(row=r, column=2, value="Итого за год, ₽").font = style.FONT_BOLD
    for col in (COL_RENT, COL_UTIL):
        letter = get_column_letter(col)
        total = ws.cell(row=r, column=col, value=f"=SUM({letter}3:{letter}{r - 2})")
        total.font = style.FONT_BOLD
        total.number_format = "# ##0.00"

    hint = ws.cell(row=r + 2, column=1,
                   value="Правки (статус, суммы, даты) сохраняются в боте: "
                         "🗂 Задачи → 📥 Загрузить правки — и пришлите этот файл. "
                         "Скрытый столбец «ID» не удаляйте.")
    hint.font = Font(italic=True)

    if last_data_row >= 3:
        status_list = _list_validation(TASK_STATUSES.values())
        ws.add_data_validation(status_list)
        status_list.add(f"E3:E{last_data_row}")

    _hide_service_column(ws, ncols)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{r - 2}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "2:2"


def _sheet_one_off(ws, conn: sqlite3.Connection) -> None:
    """Разовые задачи председателя — то, что он планирует сам."""
    ws.title = SHEET_ONE_OFF
    ncols = len(ONE_OFF_COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(row=1, column=1, value="Мои задачи (разовые)")
    title.font = style.FONT_TITLE
    title.fill = style.FILL_TITLE
    title.alignment = style.CENTER
    ws.row_dimensions[1].height = 24

    for col, name in enumerate(ONE_OFF_COLUMNS, start=1):
        ws.cell(row=2, column=col, value=name)
    style.style_header(ws, 2, ncols, ONE_OFF_WIDTHS)

    today = date.today()
    rows = repository.one_off_tasks_all(conn)
    r = 3
    for row in rows:
        v = task_service.view(row, today)
        if row["status"] in ("done", "cancelled"):
            left = ""
        elif v.days_left is None:
            left = "без срока"
        elif v.is_overdue:
            left = f"просрочено на {abs(v.days_left)} дн."
        elif v.days_left == 0:
            left = "сегодня последний день"
        else:
            left = f"осталось {v.days_left} дн."

        cells = [
            row["title"],
            task_service.category_label(row["category"]),
            _fmt(row["due_date"]),
            left,
            task_service.status_label(row["status"]),
            _fmt(row["created_at"][:10]),
            _fmt(row["done_at"]),
            row["description"],
            row["id"],
        ]
        fill = (FILL_DONE if row["status"] == "done"
                else FILL_OVERDUE if v.is_overdue
                else FILL_SOON if v.is_soon
                else FILL_ACTIVE if v.is_active_now else None)
        for col, value in enumerate(cells, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = style.BORDER
            cell.alignment = (style.LEFT if col in (1, 2, 8) else style.CENTER)
            if fill and col in (3, 4, 5):
                cell.fill = fill
        r += 1

    if not rows:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols - 1)
        hint = ws.cell(row=3, column=1,
                       value=f"{HINT_MARK} Пока пусто. Новую задачу можно вписать прямо сюда "
                             "(название, категория, срок) или поставить в боте: "
                             "🗂 Задачи → ➕ Новая задача.")
        hint.alignment = style.LEFT
        r = 4

    # Место под новые задачи: пустые строки со списками категорий и статусов
    free_rows = 15
    status_list = _list_validation(TASK_STATUSES.values())
    category_list = _list_validation(TASK_CATEGORIES.values())
    ws.add_data_validation(status_list)
    ws.add_data_validation(category_list)
    status_list.add(f"E3:E{r + free_rows}")
    category_list.add(f"B3:B{r + free_rows}")

    note = ws.cell(row=r + free_rows + 2, column=1,
                   value=f"{HINT_MARK} Новые строки вписывайте ниже последней задачи — "
                         "столбец «ID» у них останется пустым, и бот заведёт их "
                         "как новые. Чтобы сохранить правки: 🗂 Задачи → "
                         "📥 Загрузить правки и пришлите этот файл.")
    note.font = Font(italic=True)

    _hide_service_column(ws, ncols)
    ws.freeze_panes = "A3"
    if rows:
        ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{r - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "2:2"


def _sheet_verification(ws, conn: sqlite3.Connection) -> None:
    """Поверка общедомовых приборов: сроки считаются автоматически."""
    from bot.services import verification_service

    ws.title = SHEET_VERIFICATION
    ncols = len(VERIFICATION_COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(row=1, column=1,
                    value="Приборы учёта и оборудование дома")
    title.font = style.FONT_TITLE
    title.fill = style.FILL_TITLE
    title.alignment = style.CENTER
    ws.row_dimensions[1].height = 24

    for col, name in enumerate(VERIFICATION_COLUMNS, start=1):
        ws.cell(row=2, column=col, value=name)
    style.style_header(ws, 2, ncols, VERIFICATION_WIDTHS)

    today = date.today()
    rows = repository.house_meters(conn, only_active=False)
    r = 3
    for row in rows:
        v = verification_service.view(row, today)
        cells = [
            row["name"],
            "Гарантия" if v.is_warranty else "Поверка",
            row["serial"],
            _fmt(row["last_verified"]),
            row["interval_years"],
            v.next_due.strftime("%d.%m.%Y") if v.next_due else "",
            v.status_text,
            row["note"],
            row["id"],
        ]
        fill = (FILL_OVERDUE if v.is_overdue
                else FILL_SOON if (v.days_left is not None and v.days_left <= 30)
                else FILL_ACTIVE if v.is_due_soon
                else FILL_DONE if v.next_due else None)
        for col, value in enumerate(cells, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = style.BORDER
            cell.alignment = (style.LEFT if col in (1, 8) else style.CENTER)
            if fill and col in (6, 7):
                cell.fill = fill
        r += 1

    note = ws.cell(row=r + 1, column=1,
                   value="Поверка: следующая = последняя + интервал, задача "
                         "появляется за полгода до срока. Гарантия: считается "
                         "от даты ввода в эксплуатацию, напоминание — "
                         "за 3 месяца.")
    note.font = Font(italic=True)
    hint = ws.cell(row=r + 2, column=1,
                   value="Даты последней поверки и интервал можно вписать прямо "
                         "здесь и отправить файл боту: 🗂 Задачи → 📥 Загрузить "
                         "правки. Столбцы «Следующая поверка» и «Осталось» "
                         "бот пересчитает сам.")
    hint.font = Font(italic=True)

    _hide_service_column(ws, ncols)
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _sheet_rules(ws) -> None:
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 52

    ws.merge_cells("A1:D1")
    title = ws.cell(row=1, column=1,
                    value="Регламент регулярных задач председателя")
    title.font = style.FONT_TITLE
    title.fill = style.FILL_TITLE
    title.alignment = style.CENTER

    for col, name in enumerate(["Задача", "Категория", "Срок (числа)",
                                "Что делаем"], start=1):
        ws.cell(row=2, column=col, value=name)
    style.style_header(ws, 2, 4, [44, 22, 16, 52])

    for r, tpl in enumerate(DEFAULT_TASK_TEMPLATES, start=3):
        window = (f"с {tpl['day_start']} по {tpl['day_end']}"
                  if tpl["day_start"] != tpl["day_end"] else f"{tpl['day_end']}")
        if tpl["day_start"] == 1 and tpl["day_end"] < 28:
            window = f"до {tpl['day_end']}"
        cells = [tpl["title"], task_service.category_label(tpl["category"]),
                 window, tpl["description"]]
        for col, value in enumerate(cells, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = style.BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center",
                                       wrap_text=True)

    legend_row = len(DEFAULT_TASK_TEMPLATES) + 4
    ws.cell(row=legend_row, column=1, value="Подсветка в плане").font = style.FONT_BOLD
    for i, (label, fill) in enumerate([
        ("Выполнено", FILL_DONE),
        ("Срок через 3 дня и меньше", FILL_SOON),
        ("Просрочено", FILL_OVERDUE),
        ("В работе", FILL_ACTIVE),
    ], start=legend_row + 1):
        cell = ws.cell(row=i, column=1, value=label)
        cell.fill = fill
        cell.border = style.BORDER

    note_row = len(DEFAULT_TASK_TEMPLATES) + 10
    ws.cell(row=note_row, column=1,
            value="Задачи на каждый месяц создаются автоматически; цикл "
                  "продолжается в следующем году без перенастройки.").font = Font(italic=True)


def _fmt(value: str) -> str:
    if not value:
        return ""
    try:
        y, m, d = value.split("-")
        return f"{d}.{m}.{y}"
    except ValueError:
        return value
