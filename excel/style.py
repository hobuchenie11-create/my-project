"""Единый визуальный стиль DH OS для всех модулей Excel.

Цветовая схема (используется и в будущих модулях — заявки, обращения,
голосования, учет ресурсов), чтобы всё выглядело единообразно.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# --- Цвета по типу квартиры ---
FILL_1_ROOM = PatternFill("solid", fgColor="D9EAD3")   # 🟩 светло-зелёный
FILL_2_ROOM = PatternFill("solid", fgColor="CFE2F3")   # 🟦 светло-голубой
FILL_3_ROOM = PatternFill("solid", fgColor="E4D7F5")   # 🟪 светло-сиреневый
FILL_OTHER = PatternFill("solid", fgColor="FFFFFF")    # нежилые/без типа

# --- Цвета по статусу передачи ---
FILL_SUBMITTED = PatternFill("solid", fgColor="FFF2CC")     # 🟨 передали
FILL_NOT_SUBMITTED = PatternFill("solid", fgColor="F4CCCC")  # 🟥 не передали
FILL_NO_TELEGRAM = PatternFill("solid", fgColor="EFEFEF")    # 🩶 не зарегистрированы

# --- Служебное оформление ---
FILL_HEADER = PatternFill("solid", fgColor="1F4E79")     # шапка таблиц
FILL_SERVICE = PatternFill("solid", fgColor="F3F3F3")    # служебные столбцы
FILL_TITLE = PatternFill("solid", fgColor="DDEBF7")

FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
FONT_TITLE = Font(bold=True, size=14, color="1F4E79")
FONT_BOLD = Font(bold=True)

_thin = Side(style="thin", color="B7B7B7")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")

# Подписи статусов (выпадающие списки на листе «Настройки»)
STATUS_SUBMITTED = "Передал"
STATUS_NOT_SUBMITTED = "Не передал"
STATUS_NO_TELEGRAM = "Нет в Telegram"
STATUSES = [STATUS_SUBMITTED, STATUS_NOT_SUBMITTED, STATUS_NO_TELEGRAM]

ROOM_TYPES = ["1-комнатная", "2-комнатная", "3-комнатная", "Нежилое помещение"]
SOURCES = ["Telegram (бот)", "Telegram (чат)", "WhatsApp", "Вручную"]


def room_fill(rooms: int, type_: str = "residential") -> PatternFill:
    if type_ != "residential":
        return FILL_OTHER
    return {1: FILL_1_ROOM, 2: FILL_2_ROOM, 3: FILL_3_ROOM}.get(rooms, FILL_OTHER)


def room_label(rooms: int, type_: str = "residential") -> str:
    if type_ != "residential":
        return "Нежилое помещение"
    return f"{rooms}-комнатная" if rooms else "—"


def status_fill(status: str) -> PatternFill:
    return {
        STATUS_SUBMITTED: FILL_SUBMITTED,
        STATUS_NOT_SUBMITTED: FILL_NOT_SUBMITTED,
        STATUS_NO_TELEGRAM: FILL_NO_TELEGRAM,
    }.get(status, FILL_OTHER)


def style_header(ws, row: int, ncols: int, widths: list[int]) -> None:
    """Оформляет строку заголовков таблицы и задаёт ширину колонок."""
    from openpyxl.utils import get_column_letter
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = BORDER
        cell.alignment = CENTER
        if col <= len(widths):
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[row].height = 30
