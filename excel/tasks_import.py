"""Обратная загрузка правок из книги «План задач» в базу DH OS.

Excel — это выгрузка из базы, а не сама база: пока файл не загружен обратно,
бот о правках не знает. Этот модуль читает отредактированный файл и переносит
изменения в базу — после этого их видят и бот, и напоминания, и следующая
выгрузка.

Что переносится:
  • лист «Годовой план» — статус, суммы аренды и коммуналки, даты оплат,
    примечание (строки находятся по скрытому столбцу «ID»);
  • лист «Мои задачи» — правки по разовым задачам, а строки без «ID»
    заводятся как новые задачи;
  • лист «Поверка приборов» — дата последней поверки, интервал, заводской
    номер и примечание; сроки следующей поверки бот пересчитывает сам.

Пустая ячейка означает «не менять» — стереть значение через Excel нельзя,
это делается в боте. Так случайно очищенная ячейка не удаляет данные.

Запуск из терминала:  python -m excel.tasks_import reports/plan_zadach_2026.xlsx
"""
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from bot.services import task_service, verification_service
from bot.services.task_service import MONTHS_RU
from database import repository
from database.models import TASK_CATEGORIES, TASK_STATUSES
from excel.tasks_export import (COL_ID_TITLE, HINT_MARK, SHEET_ONE_OFF,
                                SHEET_PLAN, SHEET_VERIFICATION)

HEADER_ROW = 2
FIRST_DATA_ROW = 3

_STATUS_BY_LABEL = {label.lower(): code for code, label in TASK_STATUSES.items()}
_CATEGORY_BY_LABEL = {label.lower(): code for code, label in TASK_CATEGORIES.items()}


@dataclass
class ImportResult:
    """Что удалось перенести из файла в базу."""
    updated: int = 0            # исправленные задачи
    created: int = 0            # новые задачи с листа «Мои задачи»
    meters: int = 0             # обновлённые общедомовые приборы
    problems: list[str] = field(default_factory=list)

    @property
    def changed(self) -> int:
        return self.updated + self.created + self.meters

    def text(self) -> str:
        """Короткий отчёт для председателя."""
        if not self.changed and not self.problems:
            return ("Файл прочитан, но новых правок в нём нет — "
                    "в базе всё и так совпадает.")
        lines = []
        if self.updated:
            lines.append(f"✏️ Обновлено задач: {self.updated}")
        if self.created:
            lines.append(f"➕ Добавлено новых задач: {self.created}")
        if self.meters:
            lines.append(f"🔧 Обновлено приборов учёта: {self.meters}")
        if self.problems:
            lines.append("")
            lines.append("⚠️ Не удалось разобрать:")
            lines += [f"• {p}" for p in self.problems[:10]]
            if len(self.problems) > 10:
                lines.append(f"…и ещё {len(self.problems) - 10}")
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# Разбор ячеек
# ---------------------------------------------------------------------------

def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def _parse_date(value) -> str | None:
    """Дата ячейки в формате базы (ГГГГ-ММ-ДД). None — пусто или непонятно."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip().replace("/", ".").replace(" ", "")
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _parse_amount(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = (str(value).replace("\xa0", "").replace(" ", "")
            .replace("₽", "").replace(",", "."))
    try:
        return float(text)
    except ValueError:
        return None


def _parse_id(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def _header_map(ws) -> dict[str, int]:
    """Заголовок -> номер столбца. Порядок столбцов может измениться."""
    header = {}
    for col in range(1, ws.max_column + 1):
        title = _clean(ws.cell(row=HEADER_ROW, column=col).value).lower()
        if title:
            header[title] = col
    return header


def _row_values(ws, row: int, header: dict[str, int]) -> dict[str, object]:
    return {name: ws.cell(row=row, column=col).value
            for name, col in header.items()}


def _is_hint(text: str) -> bool:
    return text.startswith(HINT_MARK)


# ---------------------------------------------------------------------------
# Загрузка
# ---------------------------------------------------------------------------

def import_year_plan(conn: sqlite3.Connection, path: Path | str,
                     tg_id: int | None = None) -> ImportResult:
    """Переносит правки из файла в базу. Файл не изменяется."""
    result = ImportResult()
    wb = load_workbook(path, data_only=True)

    known = {SHEET_PLAN, SHEET_ONE_OFF, SHEET_VERIFICATION}
    if not known & set(wb.sheetnames):
        result.problems.append(
            "Это не похоже на годовой план: нет листов «Годовой план», "
            "«Мои задачи» и «Поверка приборов». Пришлите файл, выгруженный "
            "ботом (🗂 Задачи → 📊 Годовой план).")
        return result

    if SHEET_PLAN in wb.sheetnames:
        _import_tasks(conn, wb[SHEET_PLAN], SHEET_PLAN, result, tg_id)
    if SHEET_ONE_OFF in wb.sheetnames:
        _import_one_off(conn, wb[SHEET_ONE_OFF], result, tg_id)
    if SHEET_VERIFICATION in wb.sheetnames:
        _import_meters(conn, wb[SHEET_VERIFICATION], result)
        if result.meters:
            verification_service.sync_verification_tasks(conn)
    return result


def _period_from_header(text: str) -> str | None:
    """«Январь 2026» -> «2026-01». Так находится месяц в файлах без «ID»."""
    parts = text.strip().lower().split()
    if len(parts) != 2 or not parts[1].isdigit():
        return None
    if parts[0] not in MONTHS_RU:
        return None
    return f"{int(parts[1]):04d}-{MONTHS_RU.index(parts[0]) + 1:02d}"


def _task_by_title(conn: sqlite3.Connection, period: str,
                   title: str) -> object | None:
    wanted = title.strip().lower()
    for row in repository.tasks_for_period(conn, period):
        if row["title"].strip().lower() == wanted:
            return row
    return None


def _import_tasks(conn: sqlite3.Connection, ws, sheet: str,
                  result: ImportResult, tg_id: int | None) -> None:
    """Лист «Годовой план»: статусы, суммы и даты оплат.

    Строка находится по скрытому столбцу «ID». В файлах, выгруженных до его
    появления, «ID» нет — тогда задача ищется по месяцу и названию, чтобы
    заполненный файл не пришлось набивать заново.
    """
    header = _header_map(ws)
    has_id = COL_ID_TITLE.lower() in header
    period = ""

    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        values = _row_values(ws, row, header)

        # Строка-заголовок месяца: запоминаем период для поиска по названию
        month = _period_from_header(_clean(ws.cell(row=row, column=1).value))
        if month:
            period = month
            continue

        task = None
        task_id = _parse_id(values.get(COL_ID_TITLE.lower())) if has_id else None
        title = _clean(values.get("задача"))
        if task_id is not None:
            task = repository.get_task(conn, task_id)
            if task is None:
                result.problems.append(f"Лист «{sheet}», строка {row}: "
                                       f"задача №{task_id} в базе не найдена.")
                continue
        elif title and period and not _is_hint(title) and _clean(values.get("срок")):
            # Строки без даты — это итоги и подсказки, а не задачи
            task = _task_by_title(conn, period, title)
            if task is None:
                result.problems.append(
                    f"Лист «{sheet}», строка {row}: задача «{title}» за "
                    f"{task_service.period_title(period)} в базе не найдена — "
                    "строка пропущена.")
                continue
        if task is None:
            continue                      # итоги, подсказка, пустая строка

        task_id = task["id"]
        fields, changes = _task_changes(task, values, result, sheet, row,
                                        note_column=("комментарий", "примечание"),
                                        note_field="note")
        if fields:
            repository.update_task(conn, task_id, **fields)
            repository.log_task_event(conn, task_id, tg_id, "excel",
                                      "правки из Excel: " + ", ".join(changes))
            result.updated += 1


def _task_changes(task, values: dict, result: ImportResult, sheet: str,
                  row: int, note_column: str | tuple[str, ...] = "примечание",
                  note_field: str = "description") -> tuple[dict, list[str]]:
    """Сравнивает строку файла с задачей в базе. Пустая ячейка — «не менять»."""
    fields: dict[str, object] = {}
    changes: list[str] = []

    status_label = _clean(values.get("статус"))
    if status_label:
        code = _STATUS_BY_LABEL.get(status_label.lower())
        if code is None:
            result.problems.append(
                f"Лист «{sheet}», строка {row}: непонятный статус "
                f"«{status_label}». Допустимы: "
                f"{', '.join(TASK_STATUSES.values())}.")
        elif code != task["status"]:
            fields["status"] = code
            changes.append(f"статус — {TASK_STATUSES[code]}")
            if code == "done" and not task["done_at"]:
                fields["done_at"] = date.today().isoformat()

    money = [
        ("аренда, ₽", "amount", "аренда"),
        ("оплата коммуналки, ₽", "utility_amount", "коммуналка"),
    ]
    for title, column, label in money:
        amount = _parse_amount(values.get(title))
        if amount is None:
            continue
        if task[column] is None or abs(task[column] - amount) > 0.005:
            fields[column] = amount
            changes.append(f"{label} — {amount:g} ₽")

    dates = [
        ("дата поступления", "paid_at", "дата поступления аренды"),
        ("дата оплаты", "utility_paid_at", "дата оплаты коммуналки"),
        ("срок", "due_date", "срок"),
    ]
    for title, column, label in dates:
        if title not in values:
            continue
        raw = values.get(title)
        if raw is None or _clean(raw) == "":
            continue
        parsed = _parse_date(raw)
        if parsed is None:
            result.problems.append(
                f"Лист «{sheet}», строка {row}: не разобрал дату "
                f"«{_clean(raw)}» ({label}). Формат — 15.09.2026.")
        elif parsed != task[column]:
            fields[column] = parsed
            changes.append(f"{label} — {_ru(parsed)}")

    columns = (note_column,) if isinstance(note_column, str) else note_column
    note = next((_clean(values.get(c)) for c in columns if _clean(values.get(c))),
                "")
    # Описание из регламента (в старых файлах оно стояло в «Примечании»)
    # комментарием не считаем — иначе оно затрёт пустое поле.
    if note and note != task[note_field] and note != task["description"]:
        fields[note_field] = note
        changes.append(columns[0])

    return fields, changes


def _import_one_off(conn: sqlite3.Connection, ws, result: ImportResult,
                    tg_id: int | None) -> None:
    """Лист «Мои задачи»: правки и новые строки, вписанные вручную."""
    header = _header_map(ws)
    if "задача" not in header:
        result.problems.append(
            f"Лист «{SHEET_ONE_OFF}»: нет столбца «Задача».")
        return

    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        values = _row_values(ws, row, header)
        title = _clean(values.get("задача"))
        task_id = _parse_id(values.get(COL_ID_TITLE.lower()))

        if task_id is None:
            if not title or _is_hint(title):
                continue
            # Нет «ID» — задача могла быть заведена раньше (файл старой
            # выгрузки). Ищем по названию, чтобы не создать дубль.
            existing = _one_off_by_title(conn, title)
            if existing is None:
                _create_one_off(conn, title, values, result, row, tg_id)
                continue
            task_id = existing["id"]

        task = repository.get_task(conn, task_id)
        if task is None:
            result.problems.append(f"Лист «{SHEET_ONE_OFF}», строка {row}: "
                                   f"задача №{task_id} в базе не найдена.")
            continue

        fields, changes = _task_changes(task, values, result, SHEET_ONE_OFF, row)
        if title and title != task["title"]:
            fields["title"] = title
            changes.append("название")
        category = _category_code(values.get("категория"), result,
                                  SHEET_ONE_OFF, row)
        if category and category != task["category"]:
            fields["category"] = category
            changes.append(f"категория — {TASK_CATEGORIES[category]}")

        if fields:
            repository.update_task(conn, task_id, **fields)
            repository.log_task_event(conn, task_id, tg_id, "excel",
                                      "правки из Excel: " + ", ".join(changes))
            result.updated += 1


def _one_off_by_title(conn: sqlite3.Connection, title: str) -> object | None:
    wanted = title.strip().lower()
    for row in repository.one_off_tasks_all(conn):
        if row["title"].strip().lower() == wanted:
            return row
    return None


def _create_one_off(conn: sqlite3.Connection, title: str, values: dict,
                    result: ImportResult, row: int, tg_id: int | None) -> None:
    """Строка без «ID» — новая разовая задача председателя."""
    due = _parse_date(values.get("срок")) or ""
    if values.get("срок") not in (None, "") and not due:
        result.problems.append(
            f"Лист «{SHEET_ONE_OFF}», строка {row}: не разобрал срок "
            f"«{_clean(values.get('срок'))}» — задача добавлена без срока.")
    category = _category_code(values.get("категория"), result,
                              SHEET_ONE_OFF, row) or "other"
    status_label = _clean(values.get("статус"))
    status = _STATUS_BY_LABEL.get(status_label.lower(), "new")

    task_id = repository.create_task(
        conn, title, category=category, priority="normal", status=status,
        due_date=due, period=due[:7] if due else "",
        description=_clean(values.get("примечание")), source="chairman")
    repository.log_task_event(conn, task_id, tg_id, "created",
                              f"добавлена из Excel: {title}")
    result.created += 1


def _import_meters(conn: sqlite3.Connection, ws, result: ImportResult) -> None:
    """Лист «Поверка приборов»: даты поверки и межповерочные интервалы."""
    header = _header_map(ws)

    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        values = _row_values(ws, row, header)
        meter_id = _parse_id(values.get(COL_ID_TITLE.lower()))
        if meter_id is None:
            # Файл без «ID»: прибор ищем по названию
            meter = _meter_by_name(conn, _clean(values.get("прибор учёта")))
            if meter is None:
                continue
            meter_id = meter["id"]
        else:
            meter = repository.get_house_meter(conn, meter_id)
            if meter is None:
                result.problems.append(
                    f"Лист «{SHEET_VERIFICATION}», строка {row}: "
                    f"прибор №{meter_id} в базе не найден.")
                continue

        fields: dict[str, object] = {}
        serial = _clean(values.get("заводской №"))
        if serial and serial != meter["serial"]:
            fields["serial"] = serial
        note = _clean(values.get("примечание"))
        if note and note != meter["note"]:
            fields["note"] = note

        interval = values.get("интервал, лет")
        if interval not in (None, ""):
            try:
                years = int(float(interval))
            except (TypeError, ValueError):
                years = 0
            if not 1 <= years <= 20:
                result.problems.append(
                    f"Лист «{SHEET_VERIFICATION}», строка {row}: интервал "
                    f"«{_clean(interval)}» — нужно число от 1 до 20.")
            elif years != meter["interval_years"]:
                fields["interval_years"] = years

        verified = None
        raw = values.get("последняя поверка")
        if raw not in (None, "") and _clean(raw) not in ("", "—"):
            verified = _parse_date(raw)
            if verified is None:
                result.problems.append(
                    f"Лист «{SHEET_VERIFICATION}», строка {row}: не разобрал "
                    f"дату поверки «{_clean(raw)}». Формат — 15.09.2026.")

        if fields:
            repository.update_house_meter(conn, meter_id, **fields)
        if verified and verified != meter["last_verified"]:
            # Новая дата поверки — записываем в историю и пересчитываем срок
            verification_service.register_verification(
                conn, meter_id, date.fromisoformat(verified),
                note="внесено из Excel")
        elif not fields:
            continue
        result.meters += 1


def _meter_by_name(conn: sqlite3.Connection, name: str) -> object | None:
    if not name or _is_hint(name):
        return None
    wanted = name.strip().lower()
    for row in repository.house_meters(conn, only_active=False):
        if row["name"].strip().lower() == wanted:
            return row
    return None


def _category_code(value, result: ImportResult, sheet: str,
                   row: int) -> str | None:
    label = _clean(value)
    if not label:
        return None
    code = _CATEGORY_BY_LABEL.get(label.lower())
    if code is None:
        result.problems.append(
            f"Лист «{sheet}», строка {row}: непонятная категория «{label}» — "
            "выберите из выпадающего списка.")
    return code


def _ru(iso: str) -> str:
    y, m, d = iso.split("-")
    return f"{d}.{m}.{y}"


def main() -> None:
    import sys

    if len(sys.argv) < 2:
        raise SystemExit("Укажите файл: python -m excel.tasks_import "
                         "reports/plan_zadach_2026.xlsx")
    conn = repository.connect()
    try:
        result = import_year_plan(conn, sys.argv[1])
    finally:
        conn.close()
    print(result.text())


if __name__ == "__main__":
    main()
