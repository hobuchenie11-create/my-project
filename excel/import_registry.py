"""Загрузка справочника квартир из data/apartments.xlsx.

Формат файла (лист «Справочник квартир»):
    № квартиры | Тип квартиры | Кол-во ХВС | Кол-во ГВС | Эл-во

Набор приборов определяется количеством счетчиков ХВС и ГВС (они независимы):
один счетчик -> общий прибор, два -> раздельно по кухне и санузлу.
Если количество не указано, оно берется из типа квартиры: «3-комнатная» -> 2+2,
иначе 1+1. Электросчетчик у всех один.

Создать шаблон файла:      python -m excel.import_registry --template
Импортировать справочник:  python -m excel.import_registry
"""
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from bot.config import config
from database import repository
from database.models import apartment_meters, layout_label

REGISTRY_PATH = config.data_dir / "apartments.xlsx"
SHEET_NAME = "Справочник квартир"


def _to_int(value, default: int) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _rooms_from_label(label: str) -> int:
    """Число комнат из подписи типа: «3-комнатная», «2 комнатная» -> 3, 2."""
    for ch in str(label or ""):
        if ch.isdigit():
            return int(ch)
    return 0


def _counts_from_label(label: str) -> tuple[int, int]:
    """Запасной вариант, если колонки со счетчиками пусты: по типу квартиры."""
    return (2, 2) if _rooms_from_label(label) >= 3 else (1, 1)


def _pick_sheet(wb):
    return wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active


def _find_header_row(ws) -> int:
    """Номер строки с заголовками (первая колонка начинается с «№» или «номер»)."""
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        first = str(row[0] or "").strip().lower()
        if first.startswith("№") or first.startswith("номер"):
            return idx
    return 1


def _is_apartment_number(value: str) -> bool:
    text = value.strip().lower()
    return text.isdigit() or text.startswith("нежил")


def generate_template(path: Path = REGISTRY_PATH) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["№ квартиры", "Тип квартиры", "Кол-во ХВС", "Кол-во ГВС", "Эл-во"])
    for i in range(1, config.apartments_count + 1):
        ws.append([i, "1-2 комнатная", 1, 1, 1])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def import_registry(path: Path = REGISTRY_PATH,
                    conn: sqlite3.Connection | None = None) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = _pick_sheet(wb)
    header = _find_header_row(ws)

    own_conn = conn is None
    conn = conn or repository.connect()
    count = 0
    try:
        repository.create_schema(conn)
        for order, row in enumerate(
                ws.iter_rows(min_row=header + 1, values_only=True), start=1):
            if not row or row[0] is None:
                continue
            number = str(row[0]).strip()
            if not _is_apartment_number(number):
                continue
            label = row[1] if len(row) > 1 else ""
            fallback = _counts_from_label(label)
            cws = _to_int(row[2] if len(row) > 2 else None, fallback[0])
            hws = _to_int(row[3] if len(row) > 3 else None, fallback[1])

            apt_id = repository.upsert_apartment(
                conn, number, "residential", order,
                layout=layout_label(cws, hws), rooms=_rooms_from_label(label))
            repository.set_meters(conn, apt_id, apartment_meters(cws, hws))
            count += 1
        conn.commit()
    finally:
        if own_conn:
            conn.close()
    return count


if __name__ == "__main__":
    if "--template" in sys.argv:
        print(f"Шаблон справочника создан: {generate_template()}")
    else:
        print(f"Импортировано квартир: {import_registry()}")
